import pandas as pd
from openpyxl.styles import Font
from openpyxl.chart import BarChart, Reference

INPUT_CSV = "output/jobs.csv"
OUTPUT_EXCEL = "output/jobs_report.xlsx"


def auto_adjust_columns(ws):
    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter

        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            if len(value) > max_length:
                max_length = len(value)

        ws.column_dimensions[column_letter].width = max_length + 2


def get_top_companies(df):
    return (
        df["company"]
        .fillna("Unknown")
        .value_counts()
        .head(10)
        .rename_axis("company")
        .reset_index(name="jobs_count")
    )


def get_top_locations(df):
    return (
        df["location"]
        .fillna("Unknown")
        .value_counts()
        .head(10)
        .rename_axis("location")
        .reset_index(name="jobs_count")
    )


def get_top_skills(df):
    skills = (
        df["tags"]
        .dropna()
        .astype(str)
        .str.split(", ")
        .explode()
    )

    return (
        skills.value_counts()
        .head(15)
        .rename_axis("skill")
        .reset_index(name="jobs_count")
    )


def get_salary_summary(df):
    salary_df = df[["salary_min", "salary_max"]].dropna().copy()

    if salary_df.empty:
        return pd.DataFrame()

    salary_df["avg_salary"] = (
        salary_df["salary_min"] + salary_df["salary_max"]
    ) / 2

    return salary_df["avg_salary"].describe().reset_index().rename(
        columns={"index": "metric", "avg_salary": "value"}
    )


def get_overview(df):
    total_jobs = len(df)
    unique_companies = df["company"].nunique(dropna=True)
    unique_locations = df["location"].nunique(dropna=True)

    skills_count = (
        df["tags"]
        .dropna()
        .astype(str)
        .str.split(", ")
        .explode()
        .nunique()
    )

    overview = pd.DataFrame({
        "metric": [
            "Total Jobs",
            "Unique Companies",
            "Unique Locations",
            "Unique Skills"
        ],
        "value": [
            total_jobs,
            unique_companies,
            unique_locations,
            skills_count
        ]
    })

    return overview


def style_headers(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True)


def add_companies_chart(ws, max_row):
    chart = BarChart()
    chart.title = "Top Companies by Job Count"
    chart.x_axis.title = "Company"
    chart.y_axis.title = "Jobs"

    data = Reference(ws, min_col=2, min_row=1, max_row=max_row)
    categories = Reference(ws, min_col=1, min_row=2, max_row=max_row)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    ws.add_chart(chart, "D2")


def add_skills_chart(ws, max_row):
    chart = BarChart()
    chart.title = "Top Skills"
    chart.x_axis.title = "Skill"
    chart.y_axis.title = "Jobs"

    data = Reference(ws, min_col=2, min_row=1, max_row=max_row)
    categories = Reference(ws, min_col=1, min_row=2, max_row=max_row)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    ws.add_chart(chart, "D2")


def generate_report():
    df = pd.read_csv(INPUT_CSV)

    overview_df = get_overview(df)
    companies_df = get_top_companies(df)
    locations_df = get_top_locations(df)
    skills_df = get_top_skills(df)
    salary_df = get_salary_summary(df)

    with pd.ExcelWriter(OUTPUT_EXCEL, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="All Jobs", index=False)
        overview_df.to_excel(writer, sheet_name="Overview", index=False)
        companies_df.to_excel(writer, sheet_name="Top Companies", index=False)
        locations_df.to_excel(writer, sheet_name="Top Locations", index=False)
        skills_df.to_excel(writer, sheet_name="Top Skills", index=False)

        if not salary_df.empty:
            salary_df.to_excel(writer, sheet_name="Salary Summary", index=False)

        workbook = writer.book

        for sheet_name in writer.sheets:
            ws = writer.sheets[sheet_name]
            style_headers(ws)
            auto_adjust_columns(ws)

        ws_companies = writer.sheets["Top Companies"]
        add_companies_chart(ws_companies, max_row=len(companies_df) + 1)

        ws_skills = writer.sheets["Top Skills"]
        add_skills_chart(ws_skills, max_row=len(skills_df) + 1)

    print(f"Excel report saved to {OUTPUT_EXCEL}")


if __name__ == "__main__":
    generate_report()